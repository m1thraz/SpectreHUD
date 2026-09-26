"""Small Markdown-to-DOCX projection owned entirely by the DOCX plugin."""

from __future__ import annotations

import html
import re
from pathlib import Path

from docx.document import Document as DocumentType
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


_HEADING_RE = re.compile(r"^(#{1,6})\s+(.+?)\s*$")
_BULLET_RE = re.compile(r"^\s*[-*+]\s+(.+)$")
_NUMBER_RE = re.compile(r"^\s*\d+[.)]\s+(.+)$")
_IMAGE_RE = re.compile(r"^\s*!\[([^]]*)]\(([^)]+)\)\s*$")
_TABLE_SEPARATOR_RE = re.compile(r"^:?-{3,}:?$")
_HTML_TAG_RE = re.compile(r"<[^>]+>")
_LINK_RE = re.compile(r"\[([^]]+)]\([^)]+\)")
_INLINE_MARKUP_RE = re.compile(r"(?<!\\)(?:\*\*|__|~~|`|\*|_)")

_FONT_NAMES = {
    "segoe_ui": "Segoe UI",
    "calibri": "Calibri",
    "arial": "Arial",
    "lato": "Lato",
    "source_serif": "Source Serif 4",
    "georgia": "Georgia",
    "cambria": "Cambria",
}


def _plain_inline(value: str) -> str:
    value = _LINK_RE.sub(r"\1", value)
    value = _HTML_TAG_RE.sub("", value)
    value = _INLINE_MARKUP_RE.sub("", value)
    return html.unescape(value).replace("\\|", "|").strip()


def _table_cells(line: str) -> list[str]:
    return [_plain_inline(cell.strip()) for cell in line.strip().strip("|").split("|")]


def _is_table_separator(line: str) -> bool:
    cells = _table_cells(line)
    return bool(cells) and all(_TABLE_SEPARATOR_RE.fullmatch(cell.replace(" ", "")) for cell in cells)


def _shade(paragraph, fill: str) -> None:
    properties = paragraph._p.get_or_add_pPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


class DocxMarkdownRenderer:
    def __init__(self, document: DocumentType, project_dir: Path, report_font: str) -> None:
        self.document = document
        self.project_dir = project_dir.resolve()
        self.font_name = _FONT_NAMES.get(report_font, "Segoe UI")
        self.warnings: list[str] = []
        self._configure_styles()

    def _configure_styles(self) -> None:
        styles = self.document.styles
        styles["Normal"].font.name = self.font_name
        styles["Normal"].font.size = Pt(10.5)
        for level in range(1, 7):
            style = styles[f"Heading {level}"]
            style.font.name = self.font_name
        for section in self.document.sections:
            section.top_margin = Inches(0.7)
            section.bottom_margin = Inches(0.7)
            section.left_margin = Inches(0.8)
            section.right_margin = Inches(0.8)

    def render(self, markdown: str) -> tuple[str, ...]:
        lines = markdown.splitlines()
        index = 0
        while index < len(lines):
            line = lines[index]
            stripped = line.strip()
            if stripped.startswith("```"):
                index = self._code_block(lines, index)
            elif index + 1 < len(lines) and "|" in line and _is_table_separator(lines[index + 1]):
                index = self._table(lines, index)
            else:
                self._line(line)
                index += 1
        return tuple(self.warnings)

    def _code_block(self, lines: list[str], index: int) -> int:
        language = lines[index].strip()[3:].strip()
        body: list[str] = []
        index += 1
        while index < len(lines) and not lines[index].strip().startswith("```"):
            body.append(lines[index])
            index += 1
        paragraph = self.document.add_paragraph(style="No Spacing")
        run = paragraph.add_run("\n".join(body))
        run.font.name = "Consolas"
        run.font.size = Pt(8.5)
        _shade(paragraph, "F1F4F6")
        if language:
            paragraph.paragraph_format.space_before = Pt(3)
        return index + 1 if index < len(lines) else index

    def _table(self, lines: list[str], index: int) -> int:
        rows = [_table_cells(lines[index])]
        index += 2
        while index < len(lines) and "|" in lines[index] and lines[index].strip():
            rows.append(_table_cells(lines[index]))
            index += 1
        column_count = max(len(row) for row in rows)
        table = self.document.add_table(rows=len(rows), cols=column_count)
        table.style = "Table Grid"
        for row_index, values in enumerate(rows):
            for column_index, value in enumerate(values):
                cell = table.cell(row_index, column_index)
                cell.text = value
                for run in cell.paragraphs[0].runs:
                    run.font.name = self.font_name
                    run.bold = row_index == 0
        return index

    def _line(self, line: str) -> None:
        stripped = line.strip()
        if not stripped:
            return
        if stripped in {"---", "***", "___"}:
            return
        heading = _HEADING_RE.match(line)
        if heading:
            self.document.add_heading(_plain_inline(heading.group(2)), level=len(heading.group(1)))
            return
        image = _IMAGE_RE.match(line)
        if image:
            self._image(image.group(2).strip(), image.group(1).strip())
            return
        bullet = _BULLET_RE.match(line)
        if bullet:
            self.document.add_paragraph(_plain_inline(bullet.group(1)), style="List Bullet")
            return
        numbered = _NUMBER_RE.match(line)
        if numbered:
            self.document.add_paragraph(_plain_inline(numbered.group(1)), style="List Number")
            return
        if stripped.startswith(">"):
            paragraph = self.document.add_paragraph(_plain_inline(stripped.lstrip("> ")))
            paragraph.paragraph_format.left_indent = Inches(0.25)
            paragraph.runs[0].italic = True
            return
        self.document.add_paragraph(_plain_inline(stripped))

    def _image(self, raw_path: str, caption: str) -> None:
        source = Path(raw_path.strip().strip("<>"))
        if not source.is_absolute():
            source = self.project_dir / source
        try:
            resolved = source.resolve()
        except OSError:
            self.warnings.append(f"Image path could not be resolved: {raw_path}")
            return
        if not resolved.is_relative_to(self.project_dir) or not resolved.is_file():
            self.warnings.append(f"Image was skipped because it is unavailable or unsafe: {raw_path}")
            return
        try:
            self.document.add_picture(str(resolved), width=Inches(6.2))
        except (OSError, ValueError) as exc:
            self.warnings.append(f"Image could not be embedded: {raw_path} ({exc})")
            return
        if caption:
            paragraph = self.document.add_paragraph(_plain_inline(caption))
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in paragraph.runs:
                run.italic = True
                run.font.size = Pt(9)


__all__ = ["DocxMarkdownRenderer"]
