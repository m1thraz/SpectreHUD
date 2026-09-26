from pathlib import Path
import platform
import shutil
import subprocess
import sys
import zipfile

from docx import Document

from core.export_plugins import (
    ExportDataRequirement,
    PluginAvailabilityCode,
    ProjectExportContext,
    ReportExportContext,
    ReportExportRequest,
    create_bundled_export_plugin_registry,
    create_export_plugin_registry,
)
from core.reporting import ExportStatus


PLUGIN_ROOT = Path(__file__).resolve().parents[1]
REPOSITORY_ROOT = Path(__file__).resolve().parents[3]


def _loaded_plugin():
    registry = create_export_plugin_registry(external_roots=(PLUGIN_ROOT,))
    loaded = registry.load("spectrehud.docx")
    assert loaded is not None and loaded.plugin is not None
    assert loaded.availability.code is PluginAvailabilityCode.AVAILABLE
    return loaded.plugin


def test_docx_plugin_is_external_and_declares_only_required_v1_context() -> None:
    assert (
        create_bundled_export_plugin_registry().get_descriptor("spectrehud.docx") is None
    )
    registry = create_export_plugin_registry(external_roots=(PLUGIN_ROOT,))
    descriptor = registry.get_descriptor("spectrehud.docx")

    assert descriptor is not None
    assert descriptor.source_root == PLUGIN_ROOT
    assert descriptor.metadata.report_data_requirements == frozenset(
        {ExportDataRequirement.REPORT_FONT}
    )
    assert descriptor.metadata.configuration_fields == ()
    assert len(descriptor.metadata.execution_fields) == 1


def test_docx_plugin_exports_editable_structure_and_embeds_safe_images(tmp_path: Path) -> None:
    project_dir = tmp_path / "project"
    loot_dir = project_dir / "loot"
    loot_dir.mkdir(parents=True)
    shutil.copyfile(REPOSITORY_ROOT / "assets" / "spectrehud_main.png", loot_dir / "proof.png")
    destination = tmp_path / "exports"
    markdown = """\
<!-- spectre:section:start:executive_summary -->
# Security Assessment Report

## Findings Matrix

| Finding | Severity |
| --- | --- |
| Anonymous FTP | HIGH |

```text
21/tcp open ftp vsftpd 3.0.5
```

![Service proof](loot/proof.png)
<!-- spectre:section:end:executive_summary -->
"""
    request = ReportExportRequest(
        context=ReportExportContext(
            project=ProjectExportContext("IronMoth Calibration", project_dir),
            markdown=markdown,
            report_font="calibri",
        ),
        configuration={},
        execution_values={"destination": str(destination)},
    )

    result = _loaded_plugin().capabilities.report_export.export_report(request)

    assert result.status is ExportStatus.SUCCESS
    output = result.artifacts[0].path
    assert output == destination / "IronMoth_Calibration.docx"
    document = Document(output)
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert "Security Assessment Report" in text
    assert "21/tcp open ftp vsftpd 3.0.5" in text
    assert "spectre:section" not in text
    assert document.tables[0].cell(1, 0).text == "Anonymous FTP"
    assert len(document.inline_shapes) == 1


def test_docx_plugin_rejects_missing_destination_without_writing(tmp_path: Path) -> None:
    request = ReportExportRequest(
        context=ReportExportContext(
            project=ProjectExportContext("Forest", tmp_path),
            markdown="# Report",
            report_font="segoe_ui",
        ),
        configuration={},
        execution_values={},
    )

    result = _loaded_plugin().capabilities.report_export.export_report(request)

    assert result.status is ExportStatus.FAILED
    assert list(tmp_path.iterdir()) == []


def test_bundle_builder_produces_copy_installable_layout(tmp_path: Path) -> None:
    result = subprocess.run(
        [
            sys.executable,
            str(PLUGIN_ROOT / "build_bundle.py"),
            "--output-dir",
            str(tmp_path),
            "--no-vendor",
        ],
        capture_output=True,
        text=True,
        check=False,
    )

    assert result.returncode == 0, result.stderr
    archive = next(tmp_path.glob("spectrehud-docx-*.zip"))
    assert f"-{platform.system().lower()}-" in archive.name
    with zipfile.ZipFile(archive) as bundle:
        names = set(bundle.namelist())
    assert "spectrehud-docx/plugin.json" in names
    assert "spectrehud-docx/spectrehud_docx/plugin.py" in names
    assert not any("__pycache__" in name for name in names)

    install_root = tmp_path / "installed"
    with zipfile.ZipFile(archive) as bundle:
        bundle.extractall(install_root)
    code = (
        "from pathlib import Path\n"
        "from core.export_plugins import (\n"
        "    PluginAvailabilityCode, create_export_plugin_registry,\n"
        ")\n"
        f"root = Path({str(install_root)!r})\n"
        "registry = create_export_plugin_registry(external_roots=(root,))\n"
        "loaded = registry.load('spectrehud.docx')\n"
        "assert loaded is not None and loaded.plugin is not None\n"
        "assert loaded.availability.code is PluginAvailabilityCode.AVAILABLE\n"
    )
    loaded = subprocess.run(
        [sys.executable, "-c", code],
        cwd=REPOSITORY_ROOT,
        capture_output=True,
        text=True,
        check=False,
    )
    assert loaded.returncode == 0, loaded.stderr
